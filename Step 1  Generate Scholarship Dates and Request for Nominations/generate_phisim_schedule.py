#!/usr/bin/env python3
"""Generate PHISIM scholarship Excel and Word schedules from one event date.

Examples:
    python generate_phisim_schedule.py 2027-04-06
    python generate_phisim_schedule.py 04/06/2027 --template "2026 PHISIM Scholarship Process Dates.docx"

Dependencies:
    pip install openpyxl python-docx
"""

from __future__ import annotations

import argparse
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit(
        "Missing dependency. Run: pip install openpyxl python-docx"
    ) from exc


LEAD_TIMES = (67, 53, 39, 0)
MILESTONES = (
    "University qualification",
    "Student nomination information",
    "Student testimonial and photo",
    "Student recognition event",
)
NOTES = (
    "Institutions submit qualification materials; PHISIM reviews eligibility.",
    "Qualified institutions submit nominee information.",
    "Nominee submits testimonial and photo.",
    "Scholarship students are presented at the SIM Philadelphia Chapter Meeting.",
)


def parse_date(value: str) -> date:
    for pattern in ("%Y-%m-%d", "%m/%d/%Y", "%m-%d-%Y"):
        try:
            return datetime.strptime(value.strip(), pattern).date()
        except ValueError:
            pass
    raise argparse.ArgumentTypeError(
        "Use YYYY-MM-DD, MM/DD/YYYY, or MM-DD-YYYY (example: 2027-04-06)."
    )


def long_date(value: date) -> str:
    return f"{value:%A, %B} {value.day}, {value.year}"


def all_paragraphs(parent) -> Iterable:
    yield from parent.paragraphs
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from all_paragraphs(cell)


def replace_paragraph_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def generate_excel(event_date: date, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule Calculator"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:D1")
    ws["A1"] = "PHISIM Scholarship Schedule Automation"
    ws["A1"].fill = PatternFill("solid", fgColor="2F75B5")
    ws["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws["A3"] = "Annual input"
    ws.merge_cells("B3:D3")
    ws["B3"] = "Enter one date; every deadline below updates automatically."
    ws["A4"] = "Final event date"
    ws.merge_cells("B4:D4")
    ws["B4"] = event_date
    ws["B4"].number_format = "dddd, mmmm d, yyyy"
    ws["A5"] = "Instructions"
    ws.merge_cells("B5:D5")
    ws["B5"] = (
        "Change only the yellow cell above. The schedule preserves the approved "
        "67-, 53-, and 39-day lead times."
    )

    blue_fill = PatternFill("solid", fgColor="D9EAF7")
    for row in range(3, 6):
        ws.cell(row, 1).fill = blue_fill
        ws.cell(row, 1).font = Font(bold=True, color="1F1F1F")
    ws["B4"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["B4"].font = Font(size=14, bold=True, color="7F6000")
    ws["B4"].alignment = Alignment(horizontal="center")
    ws["B5"].alignment = Alignment(wrap_text=True, vertical="center")

    headers = ("Milestone", "Calculated date", "Lead time", "Notes")
    for col, value in enumerate(headers, 1):
        cell = ws.cell(8, col, value)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")

    for index, (milestone, lead, note) in enumerate(zip(MILESTONES, LEAD_TIMES, NOTES), 9):
        ws.cell(index, 1, milestone)
        ws.cell(index, 2, f"=$B$4-C{index}")
        ws.cell(index, 2).number_format = "dddd, mmmm d, yyyy"
        ws.cell(index, 3, lead)
        ws.cell(index, 3).number_format = '0 "days before"'
        ws.cell(index, 4, note)
        ws.cell(index, 1).font = Font(bold=True)
        ws.cell(index, 2).font = Font(bold=True, color="1F4E78")
        for col in range(1, 5):
            ws.cell(index, col).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[index].height = 42

    for cell in ws[12]:
        if 1 <= cell.column <= 4:
            cell.fill = PatternFill("solid", fgColor="E2F0D9")

    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows(min_row=8, max_row=12, min_col=1, max_col=4):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A15:D15")
    ws["A15"] = (
        "Rule: milestone dates are calculated by fixed lead times from the final "
        "event date. Change only B4 for each annual cycle."
    )
    ws["A15"].fill = PatternFill("solid", fgColor="F2F2F2")
    ws["A15"].font = Font(italic=True, color="595959")
    ws["A15"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[15].height = 36

    widths = (30, 29, 18, 58)
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A2"
    ws.print_area = "A1:D15"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def create_branded_fallback(event_date: date, deadlines: list[date], output_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PHILADELPHIA PHILANTHROPIC SOCIETY\nFOR INFORMATION MANAGEMENT FOUNDATION INC")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(31, 78, 121)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(f"{event_date.year} PHISIM SCHOLARSHIP PROCESS")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.7)
    table.columns[1].width = Inches(4.4)
    sections = (
        ("UNIVERSITY QUALIFICATION", deadlines[0], NOTES[0]),
        ("STUDENT NOMINATIONS", deadlines[1], NOTES[1]),
        ("STUDENT NOMINATIONS", deadlines[2], NOTES[2]),
        ("STUDENT RECOGNITION", deadlines[3], NOTES[3]),
    )
    table.rows[0].cells[0].merge(table.rows[0].cells[1])
    table.rows[0].cells[0].text = sections[0][0]
    for index, (section_name, milestone_date, note) in enumerate(sections):
        if index > 0:
            header = table.add_row().cells
            header[0].merge(header[1])
            header[0].text = section_name
        row = table.add_row().cells
        row[0].text = long_date(milestone_date)
        row[1].text = note
        row[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in row[0].paragraphs[0].runs:
            run.bold = True
        set_cell_shading(row[0], "FFF200")
    for row_index in (0, 2, 4, 6):
        cell = table.rows[row_index].cells[0]
        set_cell_shading(cell, "D9D9D9")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True

    footer = section.footer.paragraphs[0]
    footer.text = f"Revised {date.today():%Y-%m-%d}"
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.save(output_path)


def generate_word(event_date: date, deadlines: list[date], output_path: Path, template: Path | None) -> None:
    if template and template.exists():
        doc = Document(template)
        story_parts = [doc]
        for section in doc.sections:
            story_parts.extend(
                [section.header, section.first_page_header, section.even_page_header,
                 section.footer, section.first_page_footer, section.even_page_footer]
            )
        date_index = 0
        date_pattern = re.compile(
            r"^(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday), "
            r"[A-Z][a-z]+ \d{1,2}, \d{4}$"
        )
        for story in story_parts:
            for paragraph in all_paragraphs(story):
                text = paragraph.text.strip()
                if re.fullmatch(r"\d{4} PHISIM SCHOLARSHIP PROCESS", text):
                    suffix = " " if paragraph.text.endswith(" ") else ""
                    replace_paragraph_text(paragraph, f"{event_date.year} PHISIM SCHOLARSHIP PROCESS{suffix}")
                elif date_pattern.match(text) and date_index < len(deadlines):
                    replace_paragraph_text(paragraph, long_date(deadlines[date_index]))
                    date_index += 1
                elif text.startswith("Revised "):
                    replace_paragraph_text(paragraph, f"Revised {date.today():%Y-%m-%d}")
        if date_index != 4:
            raise RuntimeError(
                f"Template should contain four full milestone dates; found {date_index}."
            )
        doc.save(output_path)
    else:
        create_branded_fallback(event_date, deadlines, output_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate PHISIM Excel and Word scholarship schedules from one event date."
    )
    parser.add_argument("event_date", nargs="?", type=parse_date, help="Final event date")
    parser.add_argument(
        "--template",
        type=Path,
        default=Path("2026 PHISIM Scholarship Process Dates.docx"),
        help="Optional PHISIM Word template (default: %(default)s)",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parent,
        help="Output folder",
    )
    args = parser.parse_args()

    event_date = args.event_date
    if event_date is None:
        event_date = parse_date(input("Final event date (YYYY-MM-DD): "))

    deadlines = [event_date - timedelta(days=days) for days in LEAD_TIMES]
    output_stem = f"{event_date.year} PHISIM Scholarship Process Dates"
    excel_path = args.output_dir / f"{output_stem}.xlsx"
    word_path = args.output_dir / f"{output_stem}.docx"

    generate_excel(event_date, excel_path)
    generate_word(event_date, deadlines, word_path, args.template)

    print(f"Created: {excel_path.resolve()}")
    print(f"Created: {word_path.resolve()}")


if __name__ == "__main__":
    main()
