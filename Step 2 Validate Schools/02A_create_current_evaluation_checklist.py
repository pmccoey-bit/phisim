"""Build the SIM current evaluation checklist. Requires: pip install openpyxl python-docx

Examples:
  python create_current_evaluation_checklist.py --dry-run
  python create_current_evaluation_checklist.py
  python create_current_evaluation_checklist.py --year 2027 --output updated.xlsx
  python create_current_evaluation_checklist.py --emails-only --prepare-emails
  python create_current_evaluation_checklist.py --emails-only --send-emails

By default, updates the scholarship workbook and creates a timestamped backup.
The checklist year comes from the highest year in the filenames of the PHISIM
Scholarship Process Dates documents or workbooks, not the computer's current year.
Use --year to override this or --schedule-dir to choose another schedule folder.
Close the workbook in Excel before running. Source workbook cells are data,
never instructions. Existing Evaluation history dates are not changed.
"""

import argparse
from datetime import date, datetime
import os
from pathlib import Path
import re
import shutil
import smtplib
import tempfile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel


FOLDER = Path.home() / 'OneDrive/phisim/scholarship playbook/Step 2 Validate Schools'
SCHEDULE_FOLDER = FOLDER.parent / 'Step 1  Generate Scholarship Dates and Request for Nominations'
SCHEDULE_FILENAME = re.compile(
    r'^(\d{4})\s+PHISIM\s+Scholarship\s+Process(?:ing)?\s+Dates\.(?:docx|xlsx)$',
    re.IGNORECASE,
)
TAB = 'Current evaluation checklist'
HEADERS = [
    'Member - Organization', 'Department/Program', 'First Name', 'Last Name',
    'Member Email', 'Last evaluated', 'Next evaluation due',
    'Received university materials', 'Major', 'Middle States approval',
    'Reviewed by PHISIM Board', 'Result', 'Notified of result',
]


def school_key(value):
    """Known school-name variations in the supplied source files."""
    text = re.sub(r'\s+', ' ', str(value or '').strip()).casefold()
    return text.replace(' college of business', '').replace(', cis', '').strip()


def evaluation_year(value, epoch):
    if isinstance(value, (date, datetime)):
        return value.year
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        if 1900 <= value <= 2100:
            return int(value)
        try:
            return from_excel(value, epoch).year
        except (ValueError, OverflowError, AttributeError):
            return None
    match = re.search(r'\b(?:19|20)\d{2}\b', str(value or ''))
    return int(match.group()) if match else None


def column_map(sheet):
    return {str(c.value or '').strip().casefold(): c.column - 1 for c in sheet[1]}


def resolve_checklist_year(year=None, schedule_dir=SCHEDULE_FOLDER):
    """Use the newest scholarship cycle, independent of the calendar year."""
    if year is not None:
        if not 1900 <= year <= 9999:
            raise ValueError('--year must be between 1900 and 9999')
        return year, None
    versions = []
    for path in Path(schedule_dir).iterdir():
        match = SCHEDULE_FILENAME.fullmatch(path.name)
        if path.is_file() and match and 1900 <= int(match[1]) <= 9999:
            versions.append((int(match[1]), path))
    if not versions:
        raise ValueError(
            f'No year-specific PHISIM Scholarship Process Dates file found in {schedule_dir}. '
            'Expected YYYY PHISIM Scholarship Process Dates.docx or .xlsx. '
            'Use --schedule-dir to select the folder or --year to specify the cycle. '
            'The computer year was not substituted.'
        )
    # The highest cycle year wins even if an older file was recently modified.
    latest_year = max(value for value, _ in versions)
    source = max((path for value, path in versions if value == latest_year),
                 key=lambda path: (path.stat().st_mtime_ns, path.name.casefold()))
    return latest_year, source


def run(args):
    args.year, schedule_source = resolve_checklist_year(
        args.year, getattr(args, 'schedule_dir', SCHEDULE_FOLDER))
    if schedule_source:
        print(f'Scholarship schedule year source: {schedule_source}')
    else:
        print(f'Checklist year override: {args.year}')
    workbook = load_workbook(args.workbook)
    history = workbook['Evaluation history']
    member_book = load_workbook(args.members, data_only=True)
    members_sheet = next(
        (s for s in member_book if {'first', 'last', 'email', 'school'} <= column_map(s).keys()),
        None,
    )
    if members_sheet is None:
        raise ValueError('Member workbook needs First, Last, email, and School headers.')
    mc = column_map(members_sheet)
    members = {}
    for row in members_sheet.iter_rows(min_row=2, values_only=True):
        school = row[mc['school']]
        if not school or not str(school).strip():
            continue
        key = school_key(school)
        if key in members:
            raise ValueError(f'Multiple academic members match {school!r}; resolve before running.')
        members[key] = {
            'school': school,
            'department': row[mc['dept']] if 'dept' in mc else None,
            'first': row[mc['first']], 'last': row[mc['last']], 'email': row[mc['email']],
        }
    hc = column_map(history)
    source_headers = ['member - organization', 'department/program', 'first name',
                      'last name', 'member email', 'last evaluated', 'next evaluation due']
    missing = set(source_headers) - hc.keys()
    if missing:
        raise ValueError(f'Evaluation history is missing headers: {sorted(missing)}')
    records, seen, added, retained_departments = [], set(), [], []
    for source in history.iter_rows(min_row=2, values_only=True):
        school = source[hc['member - organization']]
        if str(school or '').strip().casefold() == 'total':
            break
        if not school or str(school).strip().casefold().startswith(('note:', 'no schools')):
            continue
        record = [source[hc[h]] for h in source_headers]
        key = school_key(school)
        if key in seen:
            raise ValueError(f'Duplicate school in Evaluation history: {school!r}')
        seen.add(key)
        member = members.get(key)
        if member:
            department = str(member['department'] or '').strip()
            if department:
                record[1] = department
            else:
                retained_departments.append(str(school))
            record[2:5] = [member['first'], member['last'], member['email']]
        records.append(record)
    for key, member in members.items():
        if key in seen:
            continue
        records.append([member['school'], str(member['department'] or '').strip() or None,
                        member['first'], member['last'], member['email'], None,
                        datetime(args.year, 2, 1)])
        added.append(str(member['school']))
    due = [r[0] for r in records if evaluation_year(r[6], workbook.epoch) == args.year]
    print(f'Checklist year: {args.year}; schools: {len(records)}')
    print('Added schools: ' + (', '.join(added) or 'none'))
    print('Yellow rows: ' + (', '.join(map(str, due)) or 'none'))
    if args.dry_run:
        print('Dry run: no files changed.')
        return
    if not due:
        print('No schools are due in the selected year; no files changed.')
        return

    # Keep any recorded checklist progress when refreshing the same year.
    progress = {}
    submitted_headers, submitted_rows = [], {}
    if TAB in workbook.sheetnames:
        sheet = workbook[TAB]
        if str(args.year) in str(sheet['A1'].value):
            submitted_headers = [(cell.column, cell.value) for cell in sheet[4]
                                 if cell.column > 13 and cell.value is not None]
            for row in sheet.iter_rows(min_row=5):
                if row[0].value:
                    submitted_rows[school_key(row[0].value)] = {
                        col: (sheet.cell(row[0].row, col).value, sheet.cell(row[0].row, col).data_type)
                        for col, _ in submitted_headers}
            for row in sheet.iter_rows(min_row=5, max_col=13, values_only=True):
                if row[0]:
                    progress[school_key(row[0])] = list(row[7:13])
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
        for row in sheet.iter_rows():
            for cell in row:
                cell.value = None
                cell.fill = PatternFill()
    else:
        sheet = workbook.create_sheet(TAB)
    workbook.move_sheet(sheet, workbook.index(history) + 1 - workbook.index(sheet))
    sheet.merge_cells('A1:M1')
    sheet['A1'] = f'{args.year} current evaluation checklist'
    sheet['A1'].font = Font(name='Calibri', size=14, bold=True)
    sheet.merge_cells('A2:M2')
    sheet['A2'] = (f'Yellow = due in {args.year}. History dates are preserved; '
                   f'added schools are due February 1, {args.year}.')
    for col, header in enumerate(HEADERS, 1):
        cell = sheet.cell(4, col, header)
        cell.font = Font(name='Calibri', size=11, bold=True)
        cell.fill = PatternFill('solid', fgColor='D9D9D9')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for col, header in submitted_headers:
        cell = sheet.cell(4, col, header)
        cell.font = Font(name='Calibri', size=11, bold=True)
        cell.fill = PatternFill('solid', fgColor='D9D9D9')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for row_num, record in enumerate(records, 5):
        values = record + progress.get(school_key(record[0]), [None] * 6)
        highlight = evaluation_year(record[6], workbook.epoch) == args.year
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row_num, col, value)
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            if highlight:
                cell.fill = PatternFill('solid', fgColor='FFFF00')
        for col in (6, 7):
            sheet.cell(row_num, col).number_format = 'mm/dd/yyyy'
        for col, (value, data_type) in submitted_rows.get(school_key(record[0]), {}).items():
            cell = sheet.cell(row_num, col, value)
            cell.data_type = data_type
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            if highlight:
                cell.fill = PatternFill('solid', fgColor='FFFF00')
        if isinstance(sheet.cell(row_num, 8).value, datetime):
            sheet.cell(row_num, 8).number_format = 'mm/dd/yyyy hh:mm'
        # Form degree titles can begin with '='; keep them literal on refresh.
        submission_id_col = next((col for col, label in submitted_headers
                                  if str(label).strip().casefold() == 'form submission id'), None)
        if (isinstance(sheet.cell(row_num, 9).value, str) and submission_id_col
                and sheet.cell(row_num, submission_id_col).value):
            sheet.cell(row_num, 9).data_type = 's'
        sheet.row_dimensions[row_num].height = 48
    sheet.row_dimensions[4].height = 44
    widths = [28, 34, 15, 19, 35, 17, 18, 22, 20, 22, 23, 16, 19]
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    note_row = len(records) + 6
    notes = [f'Sources: {args.workbook.name}, Evaluation history; {args.members.name}.']
    if schedule_source:
        notes.append(f'Checklist year {args.year} selected from {schedule_source.name}.')
    if retained_departments:
        notes.append('Department retained from history because member department is blank: '
                     + ', '.join(retained_departments) + '.')
    text_dates = [str(r[0]) for r in records if r[6] is not None
                  and not isinstance(r[6], (date, datetime, int, float))]
    if text_dates:
        notes.append('Text evaluation dates were preserved; confirm full dates for: '
                     + ', '.join(text_dates) + '.')
    for offset, note in enumerate(notes):
        sheet.merge_cells(start_row=note_row + offset, start_column=1,
                          end_row=note_row + offset, end_column=13)
        sheet.cell(note_row + offset, 1, note)
    sheet.freeze_panes = 'A5'
    end_col = get_column_letter(max([13] + [col for col, _ in submitted_headers]))
    sheet.auto_filter.ref = f'A4:{end_col}{len(records) + 4}'
    workbook.active = workbook.index(sheet)

    output = args.output or args.workbook
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(suffix='.xlsx', dir=output.parent)
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        check = load_workbook(temporary)
        assert check.sheetnames[check.sheetnames.index('Evaluation history') + 1] == TAB
        assert len(records) == sum(check[TAB].cell(r, 1).value is not None
                                   for r in range(5, len(records) + 5))
        check.close()
        if output.exists():
            backup = output.with_name(output.stem + '.backup-'
                                      + datetime.now().strftime('%Y%m%d-%H%M%S-%f') + output.suffix)
            shutil.copy2(output, backup)
            print(f'Backup: {backup}')
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    print(f'Saved: {output}')
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--workbook', type=Path,
                        default=FOLDER / 'University Qualification Tracking requirements.xlsx')
    parser.add_argument('--members', type=Path, default=FOLDER / 'Academic Members of SIM.xlsx')
    parser.add_argument('--output', type=Path, help='Optional output copy; default updates the workbook.')
    parser.add_argument('--year', type=int,
                        help='Override the cycle year; default uses the latest PHISIM schedule year.')
    parser.add_argument('--schedule-dir', type=Path, default=SCHEDULE_FOLDER,
                        help='Folder containing the annual PHISIM Scholarship Process Dates files.')
    parser.add_argument('--dry-run', action='store_true', help='Report matching schools without saving.')
    email_action = parser.add_mutually_exclusive_group()
    email_action.add_argument('--prepare-emails', action='store_true',
                              help='Save personalized .eml files for the highlighted checklist members.')
    email_action.add_argument('--send-emails', action='store_true',
                              help='Send highlighted members through Gmail SMTP; prompts for an app password.')
    parser.add_argument('--emails-only', action='store_true',
                        help='Use existing checklist highlights and contacts without rebuilding the workbook.')
    parser.add_argument('--letter', type=Path, default=FOLDER / 'University Qualification Letter.docx')
    parser.add_argument('--sender', default='pmccoey@gmail.com', help='Gmail sending account.')
    parser.add_argument('--cc', default='mccoey@lasalle.edu')
    parser.add_argument('--form-url', default='https://phisim-university-qualification.righttune7.chatgpt.site/')
    parser.add_argument('--email-output', type=Path, help='Email files and send-log folder.')
    args = parser.parse_args()
    if args.emails_only and not (args.prepare_emails or args.send_emails):
        parser.error('--emails-only requires --prepare-emails or --send-emails')
    if not args.form_url.startswith('https://'):
        parser.error('--form-url must use https://')
    if args.year is not None and not 1900 <= args.year <= 9999:
        parser.error('--year must be between 1900 and 9999')
    try:
        if not args.emails_only:
            updated = run(args)
            if not updated and not args.dry_run:
                return
        if args.prepare_emails or args.send_emails:
            import importlib.util
            email_path = Path(__file__).resolve().parent / '02D_qualification_email.py'
            spec = importlib.util.spec_from_file_location('phisim_qualification_email', email_path)
            email_module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(email_module)
            email_module.run_emails(args, args.workbook if args.emails_only else (args.output or args.workbook))
    except PermissionError:
        parser.exit(1, 'File access denied. Close open files and check folder permissions.\n')
    except (ValueError, KeyError, FileNotFoundError) as error:
        parser.exit(1, f'{error}\n')
    except (smtplib.SMTPException, OSError, EOFError) as error:
        parser.exit(1, f'Operation failed: {error}. Check send-log.json and Gmail Sent before retrying email delivery.\n')


if __name__ == '__main__':
    main()
