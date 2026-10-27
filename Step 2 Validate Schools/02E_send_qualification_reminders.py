"""Send one high-priority reminder per unanswered university, deadline + 1 day.

Uses the existing 02D letter, personalized links, attachment and delivery ledger.
No mail is sent unless --send-emails is supplied on the reminder date.
"""
import argparse
import ctypes
from datetime import date, datetime, timedelta
from getpass import getpass
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import formataddr, formatdate, make_msgid
import html
import importlib.util
import json
import os
from pathlib import Path
import re
from types import SimpleNamespace
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from docx import Document


BASE = Path(__file__).resolve().parent


def load_helper(filename):
    spec = importlib.util.spec_from_file_location(filename[:-3], BASE / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def normalize(value):
    return ' '.join(str(value or '').casefold().split())


def latest_schedule(folder, year=None):
    # Scholarship cycle year can differ from the qualification deadline year.
    pattern = re.compile(r'^(\d{4}) PHISIM Scholarship Process(?:ing)? Dates\.docx$', re.I)
    versions = [(int(m[1]), p) for p in Path(folder).iterdir() if (m := pattern.fullmatch(p.name))]
    cycle = max(y for y, _ in versions)
    if year is not None and cycle != year:
        raise ValueError('Reminder cycle changed while preparing messages.')
    source = max((p for y, p in versions if y == cycle), key=lambda p: p.stat().st_mtime_ns)
    dates = []
    for table in Document(source).tables:
        for row in table.rows:
            match = re.search(r'(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday),\s+\w+\s+\d{1,2},\s+\d{4}', row.cells[0].text)
            if match:
                dates.append(datetime.strptime(match[0], '%A, %B %d, %Y').strftime('%A, %B %d, %Y'))
    if len(dates) != 4:
        raise ValueError('Expected four dated milestones in the latest schedule.')
    return cycle, source, dates


def workbook_responses(path, year):
    """The current-cycle sheet is authoritative; old receipt dates do not count."""
    book = load_workbook(path, data_only=True)
    try:
        sheet = book['Current evaluation checklist']
        if str(year) not in str(sheet['A1'].value):
            raise ValueError('Checklist year does not match the latest schedule; no reminders sent.')
        headers = {re.sub(r'[^a-z0-9]', '', str(c.value or '').casefold()): c.column for c in sheet[4]}
        responded, seen = set(), set()
        for row in range(5, sheet.max_row + 1):
            school = normalize(sheet.cell(row, 1).value)
            if not school:
                continue
            if school in seen:
                raise ValueError('Duplicate university rows; resolve before sending reminders.')
            seen.add(school)
            def value(label):
                col = headers.get(label)
                return sheet.cell(row, col).value if col else None
            received = value('receiveduniversitymaterials')
            if isinstance(received, (datetime, date)):
                received = received.year in (year - 1, year)
            else:
                received = normalize(received) not in ('', 'no', 'n', 'false', '0', 'not received', 'pending', 'not yet', 'awaiting response')
            status = normalize(value('submissionreviewstatus'))
            if received or value('formsubmissionid') or status in ('submitted for review', 'submitted', 'received', 'approved', 'rejected', 'reviewed'):
                responded.add(school)
        return responded
    finally:
        book.close()


def pending_online_responses(form_helper, config, year):
    """Exclude durable online responses even when workbook sync is delayed."""
    base, token = form_helper.read_sync_configuration(config)
    request = Request(base + '/api/review-sync', headers={
        'Authorization': 'Bearer ' + token, 'User-Agent': 'PHISIM-Workbook-Sync/1.0'})
    with urlopen(request, timeout=30) as reply:
        items = json.load(reply)['submissions']
    if len(items) >= 50:
        raise ValueError('Online queue may be truncated. Sync all submissions before reminders.')
    return {normalize(item['payload']['university']) for item in items
            if int(item['payload']['year']) == year}


def priority_message(original, person, year, source, dates, letter, sender, cc, form_url):
    form_url = original(person, year, form_url)
    message = EmailMessage(policy=SMTP)
    message['From'], message['To'] = sender, formataddr((person['name'], person['email']))
    message['Cc'], message['Reply-To'] = cc, 'president@phisim.org'
    message['Subject'] = f'HIGH PRIORITY: Qualification reminder — {person["school"]} — {year}'
    message['Date'], message['Message-ID'] = formatdate(localtime=True), make_msgid()
    message['Importance'] = 'high'
    message['X-Priority'] = '1'
    message['X-MSMail-Priority'] = 'High'
    note = (f'We have not yet recorded a qualification response for {person["school"]}. '
            f'The qualification deadline was {dates[0]}. Please complete your submission '
            'using the personalized form below, review your responses, and select '
            'SUBMIT for Review to PHISIM. If you have already responded, please contact '
            'president@phisim.org so we can reconcile our records.')
    paragraphs = [f'Dear {person["name"] or person["school"]},', note,
        'PHISIM University Qualification Form: ' + form_url,
        f'The latest scholarship process dates, {source.name}, are attached.',
        f'Student nominations: {dates[1]}. Testimonials and photos: {dates[2]}. Scholarship recognition: {dates[3]}.',
        'Questions: president@phisim.org', 'Thank you,\nPHISIM Scholarship Committee']
    message.set_content('\n\n'.join(paragraphs))
    rich = ''.join('<p>' + html.escape(p).replace('\n', '<br>') + '</p>' for p in paragraphs)
    rich = rich.replace(html.escape(form_url), '<a href="' + html.escape(form_url, quote=True) + '">PHISIM University Qualification Form</a>')
    rich = rich.replace('Questions: president@phisim.org', 'Questions: <a href="mailto:president@phisim.org">president@phisim.org</a>')
    message.add_alternative(rich, subtype='html')
    message.add_attachment(source.read_bytes(), maintype='application', subtype='vnd.openxmlformats-officedocument.wordprocessingml.document', filename=source.name)
    return message


class Blob(ctypes.Structure):
    _fields_ = [('size', ctypes.c_ulong), ('data', ctypes.POINTER(ctypes.c_ubyte))]


def protect_password(password):
    raw = password.encode('utf-16-le')
    buffer = (ctypes.c_ubyte * len(raw)).from_buffer_copy(raw)
    incoming, outgoing = Blob(len(raw), buffer), Blob()
    if not ctypes.windll.crypt32.CryptProtectData(ctypes.byref(incoming), None, None, None, None, 0, ctypes.byref(outgoing)):
        raise OSError('Windows could not protect the credential.')
    try:
        return ctypes.string_at(outgoing.data, outgoing.size).hex()
    finally:
        ctypes.windll.kernel32.LocalFree(outgoing.data)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--send-emails', action='store_true')
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--configure-smtp', action='store_true')
    parser.add_argument('--test-reminder-date', type=date.fromisoformat,
                        help='Explicit test send date (YYYY-MM-DD); response and duplicate checks still apply.')
    args = parser.parse_args()
    mail = load_helper('02D_qualification_email.py')
    form = load_helper('02B_update_phisim_qualification_form.py')
    credentials = BASE / 'PHISIM Reminder SMTP.json'
    if args.configure_smtp:
        password = getpass('Gmail app password for pmccoey@gmail.com (hidden): ').replace(' ', '')
        if not password:
            raise ValueError('No credential entered; nothing saved.')
        credentials.write_text(json.dumps({'url': mail.FORM_URL,
            'protectedToken': protect_password(password)}, indent=2), encoding='utf-8')
        print('Credential saved with Windows user encryption. No email sent.')
        return
    schedule = BASE.parent / 'Step 1  Generate Scholarship Dates and Request for Nominations'
    mail.latest_schedule = latest_schedule
    year, source, dates = latest_schedule(schedule)
    due = datetime.strptime(dates[0], '%A, %B %d, %Y').date()
    reminder_day = args.test_reminder_date or due + timedelta(days=1)
    if args.test_reminder_date:
        print(f'Test reminder date override: {reminder_day}')
    print(f'Qualification deadline: {due}; reminder date: {reminder_day}')
    today = datetime.now(ZoneInfo('America/New_York')).date()
    if args.send_emails and not args.dry_run and today != reminder_day:
        print('Today is not the first day after the deadline. Nothing sent.')
        return
    workbook = BASE / 'University Qualification Tracking requirements.xlsx'
    responded = workbook_responses(workbook, year)
    # Fail closed if the private online queue cannot be checked.
    responded |= pending_online_responses(form, BASE / 'PHISIM Submission Sync.json', year)
    people = [p for p in mail.recipients(workbook) if normalize(p['school']) not in responded]
    print(f'{len(people)} highlighted universities have no recorded response.')
    original = mail.personalized_form_url
    mail.message_for = lambda *a: priority_message(original, *a)
    mail.recipients = lambda path: people
    options = SimpleNamespace(schedule_dir=schedule, year=year,
        letter=BASE / 'University Qualification Letter.docx', sender='pmccoey@gmail.com',
        cc='mccoey@lasalle.edu', form_url=mail.FORM_URL + f'?version=submit-for-review-{year}',
        dry_run=args.dry_run or not args.send_emails, send_emails=args.send_emails,
        email_output=BASE / f'{year} Qualification Reminder Emails')
    if args.send_emails and not options.dry_run and people:
        password = os.environ.get('PHISIM_GMAIL_APP_PASSWORD')
        if not password and credentials.exists():
            _, password = form.read_sync_configuration(credentials)
        if not password:
            raise ValueError('Gmail SMTP is not configured. Run this script with --configure-smtp locally; no reminders sent.')
        os.environ['PHISIM_GMAIL_APP_PASSWORD'] = password
    mail.run_emails(options, workbook)


if __name__ == '__main__':
    main()
