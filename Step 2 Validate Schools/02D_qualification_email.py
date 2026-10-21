"""Personalized PHISIM email preparation and opt-in Gmail SMTP delivery."""
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import formataddr, formatdate, make_msgid
from getpass import getpass
import hashlib
import html
import json
import os
from pathlib import Path
import re
import smtplib
import ssl
import base64
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode

from docx import Document
from openpyxl import load_workbook

FORM_URL = 'https://phisim-university-qualification.righttune7.chatgpt.site/'
PATTERN = re.compile(r'^(\d{4}) PHISIM Scholarship Process(?:ing)? Dates\.(docx|xlsx)$', re.I)


def latest_schedule(folder, year=None):
    versions = [(int(m[1]), p) for p in Path(folder).iterdir()
                if p.is_file() and (m := PATTERN.fullmatch(p.name))]
    if not versions:
        raise ValueError('No annual scholarship schedule found.')
    latest_year = max(y for y, _ in versions)
    if year is not None and year != latest_year:
        raise ValueError('Email year must match the latest scholarship schedule year.')
    docs = [p for y, p in versions if y == latest_year and p.suffix.lower() == '.docx']
    if not docs:
        raise ValueError('Generate the latest year Scholarship Process Dates.docx before emailing.')
    source = max(docs, key=lambda p: p.stat().st_mtime_ns)
    milestones = []
    for table in Document(source).tables:
        for row in table.rows:
            found = re.search(r'(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday),\s+\w+\s+\d{1,2},\s+\d{4}', row.cells[0].text)
            if found:
                value = datetime.strptime(found[0], '%A, %B %d, %Y')
                if value.year != latest_year:
                    raise ValueError('Schedule contains dates outside its filename year.')
                milestones.append(value.strftime('%A, %B {day}, %Y').format(day=value.day))
    if len(milestones) != 4:
        raise ValueError('Expected four dated milestones in the latest schedule document.')
    return latest_year, source, milestones


def recipients(workbook):
    book = load_workbook(workbook, data_only=True)
    try:
        sheet = book['Current evaluation checklist']
        result, seen = [], set()
        headers = {re.sub(r'[^a-z0-9]', '', str(c.value or '').casefold()): c.column for c in sheet[4] if c.value}
        for row in sheet.iter_rows(min_row=5):
            cell = row[0]
            if not (cell.fill.patternType == 'solid' and cell.fill.fgColor.type == 'rgb'
                    and cell.fill.fgColor.rgb[-6:].upper() == 'FFFF00'):
                continue
            school, department, first, last, email = [c.value for c in row[:5]]
            email = str(email or '').strip()
            if not re.fullmatch(r'[^\s@,;<>]+@[^\s@,;<>]+\.[^\s@,;<>]+', email):
                raise ValueError(f'Invalid email for {school}: {email!r}')
            if any('\r' in str(v or '') or '\n' in str(v or '') for v in (school, first, last)):
                raise ValueError(f'Unexpected line break in contact for {school}')
            key = (str(school).strip().casefold(), email.casefold())
            if key in seen:
                raise ValueError(f'Duplicate school/contact: {school}')
            seen.add(key)
            def value(label):
                col = headers.get(re.sub(r'[^a-z0-9]', '', label.casefold()))
                return sheet.cell(row[0].row, col).value if col else None
            result.append({'school': str(school), 'name': ' '.join(str(v) for v in (first, last) if v),
                           'email': email, 'department': department,
                           'saved': {field: value(label) for field, label in {
                               'school': 'Submitted school or college', 'programs': 'Major',
                               'accreditation': 'Accreditation webpage', 'faculty': 'Faculty representative',
                               'email': 'Submitted faculty email', 'fullTime': 'Full-time faculty member',
                               'member': 'Current SIM Philadelphia academic member',
                           }.items()},
                           'graduates': {label: sheet.cell(row[0].row, col).value for label, col in headers.items()
                                         if re.fullmatch(r'graduates\d{8}', label)}})
        return result
    finally:
        book.close()


def personalized_form_url(person, year, form_url):
    """Carry only this school's known form answers in the URL fragment."""
    saved = person.get('saved', {})
    payload = {'v': 1, 'year': year, 'university': person['school'],
               'school': saved.get('school') or person.get('department') or '',
               'faculty': saved.get('faculty') or person.get('name') or '',
               'email': saved.get('email') or person.get('email') or ''}
    for field in ('programs', 'accreditation', 'fullTime', 'member'):
        if saved.get(field) is not None and str(saved[field]).strip():
            payload[field] = str(saved[field])
    for index in range(3):
        value = person.get('graduates', {}).get(f'graduates{year-index-1}{year-index}')
        if isinstance(value, (int, float)) and not isinstance(value, bool) and value == int(value) and 0 <= value <= 1000000:
            payload[f'graduates{index}'] = int(value)
    encoded = base64.urlsafe_b64encode(json.dumps(payload, ensure_ascii=False, separators=(',', ':')).encode()).decode().rstrip('=')
    parsed = urlsplit(form_url)
    fragment = [(key, value) for key, value in parse_qsl(parsed.fragment) if key != 'prefill']
    fragment.append(('prefill', encoded))
    return urlunsplit(parsed._replace(fragment=urlencode(fragment)))


def message_for(person, year, source, dates, letter, sender, cc, form_url):
    form_url = personalized_form_url(person, year, form_url)
    qualification, nomination, testimonial, recognition = dates
    # The supplied Word letter is the template. Replace its placeholder paragraphs
    # with the live schedule and web form; document text never becomes commands.
    paragraphs = []
    replacements = {
        'At this time, PHISIM': f'At this time, PHISIM (Philadelphia Philanthropic Society for Information Management Foundation Inc.) is beginning its {year} student scholarship process. We are pleased to once again offer this opportunity to current juniors at qualified institutions as part of our outreach to the education community in the greater Philadelphia area. Each university is required to re-qualify every three years, and {person["school"]} is due to resubmit its qualification form.',
        'Attached to this email': f'The latest process timeline, {source.name}, is attached. Your form opens with the existing details for your school. Please review and correct them before completing the remaining answers. Complete the PHISIM University Qualification Form online: {form_url}',
        'We would like to ask': f'Please complete the PHISIM University Qualification Form by {qualification}. After it is received, the committee will review the application, and qualified universities will be sent the student nominee form.',
        'The student nominations': f'Student nominations will be requested in a separate message and are due by {nomination}. Student testimonials and photos are due by {testimonial}. The scholarship recognition ceremony will be held at the Philadelphia SIM chapter meeting on {recognition}.',
        'If you have any questions': 'If you have any questions, please email the PHISIM Board President at president@phisim.org.',
    }
    matched = set()
    for paragraph in Document(letter).paragraphs:
        text = ' '.join(paragraph.text.split())
        if not text or text in ('Process timeline', 'Criteria for University Qualification and Application Form'):
            continue
        for prefix, replacement in replacements.items():
            if text.startswith(prefix):
                text = replacement
                matched.add(prefix)
                break
        paragraphs.append(text)
    if matched != set(replacements):
        raise ValueError('The qualification letter template changed. Review its date placeholders before sending.')
    paragraphs.insert(0, f'Dear {person["name"] or person["school"]},')
    plain = '\n\n'.join(paragraphs)
    if any(int(y) != year for y in re.findall(r'\b20\d{2}\b', plain)):
        raise ValueError('The letter contains a year different from the latest schedule.')
    body = ''.join(f'<p>{html.escape(p)}</p>' for p in paragraphs)
    body = body.replace(html.escape(form_url), f'<a href="{html.escape(form_url, quote=True)}">PHISIM University Qualification Form</a>')
    body = body.replace('president@phisim.org', '<a href="mailto:president@phisim.org">president@phisim.org</a>')
    msg = EmailMessage(policy=SMTP)
    msg['From'], msg['To'] = sender, formataddr((person['name'], person['email']))
    msg['Cc'], msg['Reply-To'] = cc, 'president@phisim.org'
    msg['Subject'] = f'{year} PHISIM University Qualification — {person["school"]}'
    msg['Date'], msg['Message-ID'] = formatdate(localtime=True), make_msgid()
    msg.set_content(plain)
    msg.add_alternative(body, subtype='html')
    msg.add_attachment(source.read_bytes(), maintype='application', subtype='vnd.openxmlformats-officedocument.wordprocessingml.document', filename=source.name)
    return msg


def save_log(path, entries):
    temporary = path.with_suffix('.tmp')
    temporary.write_text(json.dumps(entries, indent=2), encoding='utf-8')
    os.replace(temporary, path)


def run_emails(args, workbook):
    year, source, dates = latest_schedule(args.schedule_dir, args.year)
    people = recipients(workbook)
    # Validate the entire batch before writing files or connecting to Gmail.
    messages = [message_for(p, year, source, dates, args.letter, args.sender, args.cc, args.form_url) for p in people]
    print(f'Email schedule: {source}; qualification deadline: {dates[0]}')
    for p in people:
        print(f'Email: {p["name"]} <{p["email"]}> — {p["school"]}')
    if args.dry_run or not people:
        print('Email dry run: nothing saved or sent.' if args.dry_run else 'No highlighted recipients; no emails sent.')
        return
    output = args.email_output or Path(workbook).parent / f'{year} Qualification Emails'
    output.mkdir(parents=True, exist_ok=True)
    ledger_path = output / 'send-log.json'
    ledger = json.loads(ledger_path.read_text(encoding='utf-8')) if ledger_path.exists() else {}
    pending = []
    for person, message in zip(people, messages):
        key = hashlib.sha256(f'{year}|{person["school"].casefold()}|{person["email"].casefold()}'.encode()).hexdigest()
        previous = ledger.get(key, {})
        if previous.get('status') == 'sent':
            print(f'Already sent; skipped: {person["school"]}')
            continue
        if args.send_emails and previous.get('status') in ('sending', 'uncertain', 'partial'):
            raise ValueError(f'Check Gmail Sent before retrying {person["school"]}; delivery in send-log.json is unresolved.')
        filename = re.sub(r'[^A-Za-z0-9]+', '-', person['school']).strip('-')
        target = output / f'{year}-PHISIM-Qualification-{filename}.eml'
        target.write_bytes(message.as_bytes())
        pending.append((key, person, message))
    print(f'Prepared {len(pending)} email files in {output}')
    if not args.send_emails or not pending:
        return
    password = os.environ.get('PHISIM_GMAIL_APP_PASSWORD') or getpass('Gmail app password (hidden): ')
    if not password.strip():
        raise ValueError('A Gmail app password is required; no messages sent.')
    # A lock also prevents simultaneous script runs from mailing the same batch.
    lock = output / '.sending.lock'
    with lock.open('x') as lock_handle:
        try:
            # Another completed run may have updated this ledger while we prepared.
            ledger = json.loads(ledger_path.read_text(encoding='utf-8')) if ledger_path.exists() else {}
            pending = [item for item in pending if ledger.get(item[0], {}).get('status') != 'sent']
            if any(ledger.get(key, {}).get('status') in ('sending', 'uncertain', 'partial') for key, _, _ in pending):
                raise ValueError('Unresolved delivery in send-log.json; check Gmail Sent before retrying.')
            with smtplib.SMTP_SSL('smtp.gmail.com', 465, context=ssl.create_default_context(), timeout=60) as smtp:
                smtp.login(args.sender, password.replace(' ', ''))
                for key, person, message in pending:
                    ledger[key] = {'school': person['school'], 'to': person['email'], 'cc': args.cc,
                                   'message_id': message['Message-ID'], 'status': 'sending',
                                   'at': datetime.now().astimezone().isoformat()}
                    save_log(ledger_path, ledger)
                    try:
                        refused = smtp.send_message(message)
                        ledger[key]['status'] = 'partial' if refused else 'sent'
                        save_log(ledger_path, ledger)
                        if refused:
                            raise ValueError(f'Some recipients were refused for {person["school"]}; check the send log.')
                    except Exception:
                        if ledger[key]['status'] == 'sending':
                            ledger[key]['status'] = 'uncertain'
                            save_log(ledger_path, ledger)
                        raise
                    print(f'Sent: {person["school"]}')
        finally:
            lock_handle.close()
            lock.unlink(missing_ok=True)
